# Read-only checks behind the manuscript's extra numbers (fanqie identity, H1c decomposition, H2 on the blind LLM coding).
# Inputs: /mnt/project-files/youwen/yisheng_dataset.csv, yisheng_claude_codes.csv, blind_coding_sheet_llm_coded.xlsx. Needs pandas, scipy, statsmodels, openpyxl.
import openpyxl
import pandas as pd, numpy as np, math
import statsmodels.api as sm, statsmodels.formula.api as smf
from scipy.stats import fisher_exact
d=pd.read_csv('/mnt/project-files/youwen/yisheng_dataset.csv',encoding='utf-8-sig',dtype=str)
A=d[d.in_analysis_set=='True']; yph=set(A[A.relation=='亦聲'].phonetic); A=A[A.phonetic.isin(yph)].copy()
A['label']=(A.relation=='亦聲').astype(int)
def gee(df,y):
    df=df.assign(y=y.astype(int))
    g=smf.gee('y ~ label',groups='phonetic',data=df,family=sm.families.Binomial(),cov_struct=sm.cov_struct.Exchangeable()).fit()
    b,se=g.params['label'],g.bse['label']; p=g.pvalues['label']
    return f'OR={math.exp(b):.2f} [{math.exp(b-1.96*se):.2f}-{math.exp(b+1.96*se):.2f}] p2={p:.2g} n={len(df)} k={df.phonetic.nunique()}'
def fis(df,y,name):
    y=y.astype(bool); a=(y&(df.label==1)).sum(); n1=(df.label==1).sum(); c=(y&(df.label==0)).sum(); n0=(df.label==0).sum()
    orr,p=fisher_exact([[a,n1-a],[c,n0-c]],alternative='greater')
    print(f'{name}: {a}/{n1} ({a/n1:.1%}) vs {c}/{n0} ({c/n0:.1%}) fisherOR={orr:.2f} p1={p:.2g} | GEE {gee(df,y)}')
mc=A[A.mc_relation!='NA'].dropna(subset=['mc_relation'])
fis(mc,mc.mc_relation=='identical','H1a')
# fanqie identity (independent of Guangyun matching)
fq=A.dropna(subset=['daxu_fanqie','head_dx_fanqie'])
fis(fq,fq.daxu_fanqie==fq.head_dx_fanqie,'Da Xu fanqie identical')
print('fq coverage', fq.label.value_counts().to_dict())
alt=mc[mc.mc_relation=='tone_voicing_alt']
fis(alt,alt.mc_qusheng_direction=='member','H1c member departing | alt')
fis(alt,alt.mc_qusheng_direction.isin(['member','head']),'departing involved | alt')
dep=alt[alt.mc_qusheng_direction.isin(['member','head'])]
fis(dep,dep.mc_qusheng_direction=='member','member departing | departing involved')
fis(mc,mc.mc_qusheng_direction.isin(['member','head']),'departing alternation among all MC pairs')
fis(mc,mc.mc_qusheng_direction=='member','member-departing among all MC pairs')
fis(mc,mc.mc_qusheng_direction=='head','head-departing among all MC pairs')
# tone/voicing breakdown for 'none'
nn=alt[alt.mc_qusheng_direction=='none']
print(pd.crosstab([nn.mc_tone_head,nn.mc_tone_member],nn.relation))
print('voicing among alt', pd.crosstab(alt.mc_voicing_differs,alt.relation))
# identity or departing member combined
fis(mc,(mc.mc_relation=='identical')|(mc.mc_qusheng_direction=='member'),'identity or member-departing')
# mc confidence
print(pd.crosstab(mc.mc_confidence,mc.relation))
print(pd.crosstab(mc.head_mc_confidence,mc.relation))

# ---- H2 on the random double-coded subsample ----
wb=openpyxl.load_workbook('/mnt/project-files/youwen/blind_coding_sheet_llm_coded.xlsx'); ws=wb['编码']
rows=list(ws.iter_rows(min_row=2,values_only=True))
s=pd.DataFrame(rows,columns=['item','series','sgloss','char','cgloss','code','conf','note'])
c=pd.read_csv('/mnt/project-files/youwen/yisheng_claude_codes.csv',encoding='utf-8-sig',dtype=str)
s['item']=s['item'].astype(str); m=s.merge(c,on='item',suffixes=('','_c'))
assert (m.char==m.char_c).all()
d=pd.read_csv('/mnt/project-files/youwen/yisheng_dataset.csv',encoding='utf-8-sig',dtype=str)
m=m.merge(d[['phonetic','char','relation','in_analysis_set','paronomastic_gloss','mc_relation','mc_qusheng_direction','is_xinfu']],left_on=['series','char'],right_on=['phonetic','char'],how='left',suffixes=('','_d'))
m=m[(m.relation_d==m.relation)|m.relation_d.isna()|~m.item.duplicated(keep=False)]
print(len(m));print(m.relation_d.value_counts(dropna=False).to_dict(), m.in_analysis_set.value_counts(dropna=False).to_dict())
print(m[m.relation!=m.relation_d][['item','series','char','relation','relation_d']])
a=m[m.in_analysis_set=='True'].copy(); a['relation']=a['relation_d']
print('analysis-set items',len(a),a.relation.value_counts().to_dict())
def rep(df,col,name):
    df=df[df[col]!='X'].copy(); df['y']=df[col].isin(['Y','E']).astype(int); df['label']=(df.relation=='亦聲').astype(int)
    y=df[df.label==1]; o=df[df.label==0]
    orr,p=fisher_exact([[y.y.sum(),len(y)-y.y.sum()],[o.y.sum(),len(o)-o.y.sum()]],alternative='greater')
    ys=(y[col]=='Y').sum(); os_=(o[col]=='Y').sum()
    orr2,p2=fisher_exact([[ys,len(y)-ys],[os_,len(o)-os_]],alternative='greater')
    print(f'{name}: Y+E {y.y.sum()}/{len(y)} vs {o.y.sum()}/{len(o)} OR={orr:.1f} p={p:.2g} | Y only {ys}/{len(y)} vs {os_}/{len(o)} OR={orr2:.1f} p={p2:.2g}')
    return df
for col,nm in [('code','LLM blind'),('code_core','Claude first-pass')]:
    rep(a,col,nm+' all')
    rep(a[a.paronomastic_gloss=='False'],col,nm+' no paronomastic')
# cross semantic x MC
b=a[a.code!='X'].copy(); b['semrel']=b.code.isin(['Y','E']).astype(int); b['ident']=(b.mc_relation=='identical').astype(int); b['label']=(b.relation=='亦聲').astype(int)
b=b[b.mc_relation.notna()&(b.mc_relation!='NA')]
print(pd.crosstab([b.relation,b.semrel],b.mc_relation))
g=smf.logit('ident ~ label + semrel',data=b).fit(disp=0); print(g.summary2().tables[1][['Coef.','P>|z|']])
g=smf.logit('semrel ~ label + ident',data=b).fit(disp=0); print(g.summary2().tables[1][['Coef.','P>|z|']])
print('phonetics in sample', a.series.nunique())
print('paronomastic among labelled sample', pd.crosstab(a.relation,a.paronomastic_gloss))
df=a[a.code!='X'].copy(); df['y']=df.code.isin(['Y','E']).astype(int); df['label']=(df.relation=='亦聲').astype(int)
g=smf.gee('y ~ label',groups='series',data=df,family=sm.families.Binomial(),cov_struct=sm.cov_struct.Exchangeable()).fit()
b_,se=g.params['label'],g.bse['label']; print('H2 GEE',len(df),df.series.nunique(),math.exp(b_),math.exp(b_-1.96*se),math.exp(b_+1.96*se),g.pvalues['label']/2)
df2=df[df.paronomastic_gloss=='False']
g=smf.gee('y ~ label',groups='series',data=df2,family=sm.families.Binomial(),cov_struct=sm.cov_struct.Exchangeable()).fit()
b_,se=g.params['label'],g.bse['label']; print('H2 GEE nopar',len(df2),math.exp(b_),math.exp(b_-1.96*se),math.exp(b_+1.96*se),g.pvalues['label']/2)
print('n in identity-semantic model', len(b))
ps=[2.9e-07,0.1,0.31,0.0019,g.pvalues['label']/2]

# The sensitivity analysis without Wang Yun's rejected labels is the pilot thread's: scripts/wangyun_sensitivity.py -> youwen/yisheng_models_wangyun_sensitivity.csv (used in Tables 2, 4 and 5).
