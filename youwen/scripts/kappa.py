# 用法：python3 kappa.py <已填写的 blind_coding_sheet.xlsx> <youwen_pilot.csv>
# 输出 Cohen's κ（四类 Y/E/N/X；以及严口径 Y vs 非Y）与分歧清单
import sys,csv
from collections import Counter
from openpyxl import load_workbook
ws=load_workbook(sys.argv[1])['编码']
B={}
for r in ws.iter_rows(min_row=2,values_only=True):
    if r[7]: B[(r[1],r[3])]=str(r[7]).strip().upper()
A={(r['series'],r['char']):r['code_core'] for r in csv.DictReader(open(sys.argv[2],encoding='utf-8-sig'))}
pairs=[(A[k],b) for k,b in B.items() if k in A]
def kappa(p):
    n=len(p); po=sum(a==b for a,b in p)/n
    ca=Counter(a for a,_ in p); cb=Counter(b for _,b in p)
    pe=sum(ca[k]*cb[k] for k in set(ca)|set(cb))/n/n
    return (po-pe)/(1-pe), po
k4,po4=kappa(pairs); k2,po2=kappa([(a=='Y',b=='Y') for a,b in pairs])
print(f'n={len(pairs)}  四类 κ={k4:.3f} (一致率 {po4:.1%})  Y/非Y κ={k2:.3f} (一致率 {po2:.1%})')
for k,b in B.items():
    if k in A and A[k]!=b: print('分歧',k,'Claude:',A[k],'你:',b)
