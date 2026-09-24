# 生成第二编码人盲编码表：各系列（含对照组）分层随机抽 30%，打乱顺序，隐藏组别与 Claude 编码
import csv,math,random,sys,os
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font,Alignment
out=sys.argv[1]
rows=[r for r in csv.DictReader(open(os.path.join(out,'youwen_pilot.csv'),encoding='utf-8-sig')) if r['code_core']!='H']
rnd=random.Random(20260924)
by={}
for r in rows: by.setdefault(r['series'],[]).append(r)
samp=[]
for s,L in by.items(): samp+=rnd.sample(L,math.ceil(0.3*len(L)))
rnd.shuffle(samp)
wb=Workbook(); ws=wb.active; ws.title='编码'
H=['item','声符','预设核心义','字','说文释义','大徐反切','段注相关语句','编码(Y/E/N/X)','把握(1-3)','备注']
ws.append(H)
for i,r in enumerate(samp,1):
    ws.append([i,r['series'],r['claimed_core'],r['char'],r['shuowen_gloss'],r['daxu_fanqie'],r['duan_semantic_remarks'],'','',''])
dv=DataValidation(type='list',formula1='"Y,E,N,X"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'H2:H{len(samp)+1}')
dv2=DataValidation(type='list',formula1='"1,2,3"',allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'I2:I{len(samp)+1}')
for col,w in zip('ABCDEFGHIJ',[6,6,16,6,40,9,40,12,9,24]): ws.column_dimensions[col].width=w
for c in ws[1]: c.font=Font(bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2'
g=wb.create_sheet('说明',0)
for line in ['盲编码说明（请勿查看 youwen_pilot.csv 的 code_core / coder_note 列）','',
 '任务：判断每个字的《说文》本义是否含"预设核心义"。',
 'Y = 释义本身直接含核心义或其近义词',
 'E = 只能借助段注"之言/引伸"、后世训诂或一步以上推理才连得上',
 'N = 本义与核心义无关',
 'X = 专名（地名、水名、草木鸟兽虫鱼名等），语义不透明',
 '把握：1 = 不确定，2 = 较确定，3 = 很确定',
 '','规则：以本义为准，不以后起义、假借义为准；核心义已事先给定，不要自己另定；可查《说文》原书、段注，但不要看 Claude 的编码。',
 '表中各声符的条目已打乱顺序；其中有些声符是对照组，不告诉你是哪几个。',
 f'共 {len(samp)} 条，按各声符 30% 分层随机抽样（随机种子 20260924）。']:
    g.append([line])
g.column_dimensions['A'].width=110
wb.save(os.path.join(out,'blind_coding_sheet.xlsx'))
with open(os.path.join(out,'blind_coding_key.csv'),'w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(['item','row_id','series','char']); [w.writerow([i,r['row_id'],r['series'],r['char']]) for i,r in enumerate(samp,1)]
print(len(samp))
