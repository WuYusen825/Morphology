#!/usr/bin/env python3
"""Checks the submission files of a paper version against what Morphology's submission guidelines and the Snapp page ask for.

It reads youwen/manuscript/yisheng_paper_<v>.md and _anonymised.md, yisheng_<v>_title_page.md, yisheng_<v>_cover_letter.md, the
Word files in youwen/manuscript/submission/ and the figure files Fig1.*, and prints one line per check: PASS, FAIL or INFO
(for what only a person can judge or what is only reported). The exit status is 1 if any check fails. The checks cover
only what a script can see; the checklist in submission/SUBMISSION_CHECKLIST.md says which requirements need the authors.

If youwen/manuscript/submission/supplement/ exists, the seven Online Resource files in it are checked against the captions of the
anonymised manuscript as well (names, captions, row and file counts, empty properties, no names, e-mail addresses or local paths).
With --rerun-esm7 the analysis package ESM_7.zip is unpacked into a temporary folder and run (about a minute).

Run from anywhere:  python3 youwen/scripts/yisheng_submission_audit.py v10 [--rerun-esm7]
"""
import io
import re
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

M = Path(__file__).resolve().parents[1] / "manuscript"
SUB = M / "submission"
REVEALING = ["Yusen", "Weiyi", "WuYusen", "Beijing Foreign", "BFSU", "Foreign Studies", "github.com/WuYusen825"]
UNNUMBERED = {"References", "Supplementary Information", "Statements and Declarations"}
results = []


def check(ok, name, detail="", info=False):
    results.append(("INFO" if info else ("PASS" if ok else "FAIL"), name, detail))


def plain(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()


EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
LOCAL_PATH = re.compile(r"/home/|/mnt/|/tmp/|/root/|/Users/|[A-Z]:\\Users")
SESSION_ID = re.compile(r"claude\.ai/|session_[0-9A-Za-z]{10,}|cmsg_|cse_")
SUPPLEMENT_REVEALING = REVEALING + ["Wu Yusen", "Qu Weiyi", "English and International Studies"]


def norm(s):
    """Caption text compared without Markdown emphasis, line breaks and doubled blanks."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", s.replace("*", ""))).strip()


def text_of_member(name, data):
    if name.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        parts = [" ".join(wb.sheetnames)]
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        parts.append(c.value)
                    if c.comment:
                        parts.append(c.comment.text + (c.comment.author or ""))
        return "\n".join(parts)
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def xlsx_props(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    core = z.read("docProps/core.xml").decode("utf-8") if "docProps/core.xml" in z.namelist() else ""
    app = z.read("docProps/app.xml").decode("utf-8") if "docProps/app.xml" in z.namelist() else ""
    filled = [t for t in ("creator", "lastModifiedBy", "description", "subject", "keywords", "category")
              if re.search(rf"<(?:dc|cp):{t}>[^<]+</", core)]
    filled += [t for t in ("Company", "Manager") if re.search(rf"<{t}>[^<]+</", app)]
    return filled


def supplement_checks(anon, rerun):
    sup = SUB / "supplement"
    if not sup.is_dir():
        check(True, "supplement files (submission/supplement/)", "not in the repository yet", info=True)
        return
    import openpyxl
    import pymupdf
    listed = {int(n): (f, norm(cap)) for n, f, cap in re.findall(r"^\*\*Online Resource (\d+)\*\* \((ESM_\d+\.\w+)\) (.*)$", anon, re.M)}
    present = sorted(p.name for p in sup.iterdir())
    check(present == sorted(f for f, _ in listed.values()) and len(listed) == 7,
          "supplement folder holds exactly the files named in the manuscript, ESM_n.ext", f"{present}")
    blobs = {}
    for n, (fname, cap) in sorted(listed.items()):
        path = sup / fname
        if not path.exists():
            check(False, f"{fname} exists")
            continue
        data = path.read_bytes()
        blobs[fname] = data
        # the caption inside the file (title block of the PDF, sheet About, README of the zip) is the manuscript's caption
        if fname.endswith(".pdf"):
            d = pymupdf.open(stream=data, filetype="pdf")
            first = norm(d[0].get_text())
            check(cap in first and "Withheld for double-anonymous review" in first and bool(re.search(r"Journal:? Morphology", first)),
                  f"{fname}: caption on page 1 identical to the manuscript's, authors and corresponding author withheld", f"{d.page_count} pages")
            md = d.metadata
            filled = [k for k in ("author", "subject", "keywords", "creator", "producer") if md.get(k)]
            check(not filled, f"{fname}: PDF properties without author, creator or producer", f"title: {md.get('title')!r}; filled: {filled}")
        elif fname.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
            about = {str(r[0]): r[1] for r in wb["About"].iter_rows(values_only=True) if r and r[0]}
            check(norm(str(about.get("Caption", ""))) == cap and about.get("Authors") == about.get("Corresponding author (affiliation, e-mail)") == "Withheld for double-anonymous review"
                  and about.get("Journal") == "Morphology",
                  f"{fname}: sheet About carries the manuscript's caption, authors and corresponding author withheld",
                  f"sheets {len(wb.sheetnames)}, first {wb.sheetnames[0]}, last {wb.sheetnames[-1]}")
            check(wb.sheetnames[0] == "About" and wb.sheetnames[-1] == "data_dictionary", f"{fname}: first sheet About, last sheet data_dictionary")
            check(not xlsx_props(data), f"{fname}: workbook properties empty (no creator, last modifier, company)", f"filled: {xlsx_props(data)}")
        else:
            z = zipfile.ZipFile(io.BytesIO(data))
            readme = [m for m in z.namelist() if m.endswith("README.txt") and m.count("/") == 1]
            txt = norm(z.read(readme[0]).decode("utf-8-sig")) if readme else ""
            check(bool(readme) and ("Caption: " + cap) in txt and "Withheld for double-anonymous review" in txt,
                  f"{fname}: README caption identical to the manuscript's, authors and corresponding author withheld", f"{len(z.namelist())} members")
            roots = {m.split("/")[0] for m in z.namelist()}
            check(len(roots) == 1 and not any(re.search(r"(^|/)(\.|__MACOSX|Thumbs\.db)", m) for m in z.namelist()) and z.comment == b"",
                  f"{fname}: one root folder, no hidden or system files, no zip comment", f"{sorted(roots)}")
            bad = [m for m in z.namelist() if m.endswith(".xlsx") and xlsx_props(z.read(m))]
            check(not bad, f"{fname}: workbooks inside have empty properties", f"{bad}")
        # nothing that identifies the authors, in the text of any file or member
        members = {fname: data}
        if fname.endswith(".zip"):
            members = {f"{fname}/{m}": z.read(m) for m in z.namelist() if not m.endswith("/")}
        hits = []
        for label, blob in members.items():
            if label.endswith(".pdf"):
                d = pymupdf.open(stream=blob, filetype="pdf")
                text = "\n".join(pg.get_text() for pg in d) + str(d.metadata)
            elif label.endswith((".png", ".jpg", ".eps", ".tif", ".tiff", ".pdf")):
                continue
            else:
                text = text_of_member(label, blob)
            # self_check.py lists the path prefixes it looks for ('/home/', '/tmp/'); that list is not a hit
            paths = [] if label.endswith("scripts/self_check.py") else [LOCAL_PATH]
            for pat in [re.compile(re.escape(b), re.I) for b in SUPPLEMENT_REVEALING] + [EMAIL, SESSION_ID] + paths:
                m = pat.search(text)
                if m:
                    hits.append(f"{label}: {m.group(0)!r}")
        check(not hits, f"{fname}: no author name, affiliation, e-mail address, local path or session id in any text or member", f"{hits[:4]}")

    # row and file counts that the captions and the text state
    if "ESM_2.xlsx" in blobs and "ESM_3.xlsx" in blobs:
        rows = {}
        for fname, sheets in (("ESM_2.xlsx", ("pairs", "recension_collation", "daxu_spotcheck")),
                              ("ESM_3.xlsx", ("first_sample", "enlarged_coding", "author_check", "second_answers"))):
            wb = openpyxl.load_workbook(io.BytesIO(blobs[fname]), read_only=True)
            for s in sheets:
                rows[s] = sum(1 for r in wb[s].iter_rows(min_row=2, values_only=True) if any(c is not None for c in r))
        want = {"pairs": 1333, "recension_collation": 227, "daxu_spotcheck": 104, "first_sample": 100, "enlarged_coding": 767,
                "author_check": 126, "second_answers": 203}
        check(rows == want, "data rows per sheet (1,333 pairs; 227 collated; 104 checked; 100, 767, 126 and 203 codings)", f"{rows}")
        stated = {"1,333 extracted pairs", "767 pairs", "677 new items", "90 first-sample items", "126 pairs", "227 labelled entries", "104 entries"}
        check(all(x in anon for x in stated), "the numbers behind these row counts are the ones stated in the captions", f"{sorted(stated)}")
    if "ESM_5.zip" in blobs:
        names = zipfile.ZipFile(io.BytesIO(blobs["ESM_5.zip"])).namelist()
        pr = [m for m in names if "/prompts/" in m and m.endswith(".txt")]
        ra = [m for m in names if "/raw_answers/" in m and m.endswith(".txt")]
        check(len(pr) == 14 and len(ra) == 16 and sum("second_answer_not_used" in m for m in ra) == 2,
              "ESM_5: 14 prompts, 14 raw answers and the 2 unused second answers", f"{len(names)} files: {len(pr)} prompts, {len(ra)} answer files")
    if "ESM_7.zip" in blobs:
        names = zipfile.ZipFile(io.BytesIO(blobs["ESM_7.zip"])).namelist()
        check(len(names) == 69, "ESM_7: 69 members (67 files of the package, the Apache-2.0 licence text and the NOTICE)", f"{len(names)}")
        if rerun:
            with tempfile.TemporaryDirectory() as tmp:
                zipfile.ZipFile(io.BytesIO(blobs["ESM_7.zip"])).extractall(tmp)
                root = next(Path(tmp).iterdir())
                done = subprocess.run([sys.executable, "scripts/run_all.py"], cwd=root, capture_output=True, text=True, timeout=1200)
                m = re.search(r"(\d+) checks passed, (\d+) failed", done.stdout)
                check(done.returncode == 0 and bool(m) and m.group(2) == "0",
                      "ESM_7: unpacked and run (python scripts/run_all.py): all checks pass, exit status 0", m.group(0) if m else done.stdout[-300:])
        else:
            check(True, "ESM_7 rerun", "skipped (add --rerun-esm7); last result in the checklist, Section 8", info=True)


def main(v, rerun=False):
    full = (M / f"yisheng_paper_{v}.md").read_text(encoding="utf-8")
    anon = (M / f"yisheng_paper_{v}_anonymised.md").read_text(encoding="utf-8")
    page = (M / f"yisheng_{v}_title_page.md").read_text(encoding="utf-8")
    letter = (M / f"yisheng_{v}_cover_letter.md").read_text(encoding="utf-8")

    # ---- abstract and keywords ------------------------------------------------------------------------------------
    ab = re.sub(r"[*\\]", "", anon[anon.index("**Abstract**") + 12:anon.index("**Keywords")].strip())
    ws = len(ab.split())
    cjk = len(re.findall(r"[㐀-鿿\U00020000-\U0002fa1f]", ab))
    word_like = ws + max(cjk - len(re.findall(r"[㐀-鿿\U00020000-\U0002fa1f]+", ab)), 0)
    rx = len(re.findall(r"\w+", ab))
    check(150 <= ws <= 250 and 150 <= word_like <= 250, "abstract 150-250 words",
          f"{ws} words by whitespace, {word_like} if each Chinese character counts, {rx} by \\w+ (decimals split)")
    check(rx <= 255, "abstract margin", f"\\w+ count {rx}", info=True)
    kw = re.search(r"^\*\*Keywords\*\* (.*)$", anon, re.M).group(1).split(" · ")
    check(4 <= len(kw) <= 6, "4-6 keywords", f"{len(kw)}: {'; '.join(kw)}")
    check(not re.search(r"\b(?:MC|OC|OR|CI|GEE|MH)\b", ab), "no abbreviations in the abstract")

    # ---- headings ------------------------------------------------------------------------------------------------
    heads = re.findall(r"^(#{2,6}) (.*)$", anon, re.M)
    bad = [h for lvl, h in heads if not (
        (len(lvl) == 2 and (re.match(r"\d+ \S", h) or h in UNNUMBERED)) or (len(lvl) == 3 and re.match(r"\d+\.\d+ \S", h)))]
    check(not bad, "decimal headings, at most three levels", f"{len(heads)} headings; odd ones: {bad}")
    subs = [h for lvl, h in heads if len(lvl) == 3]
    check(any(h.startswith("3.6 Use of large language models") for h in subs), "LLM use documented in the Methods (Section 3.6)")

    # ---- tables and figure -------------------------------------------------------------------------------------------
    body_lines = anon.split("\n")
    caps = [int(m.group(1)) for ln in body_lines if (m := re.match(r"\*\*Table (\d+)\*\* ", ln))]
    check(caps == list(range(1, len(caps) + 1)), "tables numbered 1..n with Arabic numerals", f"{len(caps)} tables")
    first = {}
    for i, ln in enumerate(body_lines):
        if ln.startswith(("**Table", "|", "*Note.*")):
            continue
        for m in re.finditer(r"\bTable (\d+)\b", ln):
            first.setdefault(int(m.group(1)), i)
    order = sorted(first, key=first.get)
    check(order == sorted(first) and set(first) == set(caps), "tables cited in text in consecutive numerical order",
          f"first citations in this order: {order}")
    for n in caps:
        cap_i = next(i for i, ln in enumerate(body_lines) if ln.startswith(f"**Table {n}** "))
        j = cap_i + 1
        while body_lines[j].strip() == "":
            j += 1
        assert body_lines[j].startswith("|"), f"Table {n}: the table does not follow its caption"
        k = j
        while body_lines[k].startswith("|"):
            k += 1
        cells = "\n".join(body_lines[j:k])
        note = body_lines[k + 1] if body_lines[k].strip() == "" else body_lines[k]
        marks_cells = set(re.findall(r"\^([a-z])\^", cells))
        marks_note = set(re.findall(r"\^([a-z])\^", note)) if note.startswith("*Note.*") else set()
        check(marks_cells == marks_note, f"Table {n}: footnote letters in the body match the note",
              f"body {sorted(marks_cells)}, note {sorted(marks_note)}")
        check(not body_lines[cap_i].rstrip().endswith((".", ":")) or True, f"Table {n}: caption", "", info=True)
    figs = [ln for ln in body_lines if ln.startswith("![")]
    check(len(figs) == 1 and figs[0].startswith("![**Fig. 1** "), "figure caption begins with bold 'Fig. 1'")
    cap_text = re.match(r"!\[(.*)\]\(.*\)", figs[0]).group(1)
    check(not re.match(r"\*\*Fig\. 1\*\*[.:,;]", cap_text) and not cap_text.rstrip().endswith((".", ":", ";", ",")),
          "no punctuation after the figure number or at the end of the caption")
    check("Figure" not in re.sub(r"!\[.*\]\(.*\)", "", anon.split("## References")[0]) and "Fig. 1" in anon,
          "figure cited as 'Fig. 1' in the text, never 'Figure 1'")
    check("(submission/Fig1.png){width=119mm}" in figs[0], "figure embedded from Fig1.png at 119 mm")
    check("no colour" not in cap_text and not re.search(r"\b(red|blue|green|yellow|colou?r)\b", cap_text, re.I),
          "caption does not rely on colour words")

    # ---- references ---------------------------------------------------------------------------------------------
    refs = [l for l in anon.split("## References")[1].split("\n") if l.strip() and not l.startswith("#")]
    def year_key(r):
        # "(2005)", "(100 CE)" and "(10th century)" all sort by date; an entry without a date sorts last
        m = re.search(r"\((?:ca\. )?(\d{1,4})(?: CE|th century)?[,)]", r) or re.search(r"\((?:ca\. )?(\d{1,4})", r)
        if not m:
            return 9999
        return int(m.group(1)) * (100 if "th century" in m.group(0) + r[m.end():m.end() + 12] else 1)

    keys = [(re.sub(r"[^a-z]", "", plain(re.sub(r"\[[^\]]*\]", "", r.split(" (")[0]))), year_key(r)) for r in refs]
    check(keys == sorted(keys), "reference list alphabetised by first author, then year", f"{len(refs)} entries")
    no_doi = [r[:40] for r in refs if "https://doi.org/" not in r]
    check(all("doi:" not in r.lower().replace("https://doi.org/", "") for r in refs), "DOIs written as full https://doi.org/ links")
    check(True, "entries with a DOI", f"{len(refs) - len(no_doi)} of {len(refs)}; without: {len(no_doi)} (books, Chinese-language works and proceedings Crossref does not know)", info=True)
    check(all(re.search(r"\*[^*]+\*", r) or "Cited from" in r for r in refs), "journal and book titles italicised")
    text_only = re.sub(r"^\|.*$", "", anon.split("## References")[0], flags=re.M)
    missing = []
    for r in refs:
        sur = re.sub(r"\[.*?\]", "", r.split(",")[0]).strip()
        ym = re.search(r"\((?:ca\. )?(\d{4})", r)
        if not ym:
            continue
        yr = ym.group(1)
        orig = re.search(r"Original work (?:published|completed) (?:ca\. )?(\d{3,4})", r)
        pat = rf"{re.escape(sur)}[^()]{{0,60}}?(?:\(|, | )(?:ca\. )?(?:{yr}|[0-9]{{3,4}}/{yr}|{yr}/[0-9]{{4}}|{re.escape(orig.group(1)) if orig else yr})"
        if not re.search(pat, text_only) and not re.search(rf"{re.escape(sur)}.{{0,80}}{yr}", text_only):
            missing.append(r[:50])
    check(not missing, "every dated reference is cited in the text", f"not found: {missing}")
    cited = set(re.findall(r"([A-Z][A-Za-zÀ-ɏ'’\-]+)(?: et al\.|(?: &| and) [A-Z][A-Za-zÀ-ɏ'’\-]+)?,? \(?(?:ca\. )?((?:1[0-9]|20)\d\d)(?:/\d{4})?\)?", text_only))
    refsy = {(re.sub(r"\[.*?\]", "", r.split(",")[0]).strip(), re.search(r"\((?:ca\. )?(\d{4})", r).group(1)) for r in refs if re.search(r"\((?:ca\. )?(\d{4})", r)}
    stray = sorted(f"{a} {y}" for a, y in cited if (a, y) not in refsy and not any(a == s and y in (yy, yy) for s, yy in refsy))
    check(True, "in-text name-year pairs without an entry (for a person to look at)", "; ".join(stray[:40]) or "none", info=True)

    # ---- statements, online resources -----------------------------------------------------------------------------------
    # Snapp's double-anonymous instructions (springernature.com/gp/snapp/submitting/how-to-submit/double-anonymous): the manuscript
    # file carries no acknowledgement, contribution, competing-interest, ethics or funding statement (the system asks for them, and
    # the title page file is their source text); the data availability statement is not on that list and stays.
    sd = anon[anon.index("## Statements and Declarations"):anon.index("## References")]
    check("**Data availability**" in sd, "Statements and Declarations in the manuscript: Data availability")
    for label in ("Funding", "Competing interests", "Ethics approval and consent", "Author contributions", "Acknowledgements", "Acknowledgments"):
        check(f"**{label}**" not in anon and f"## {label}" not in anon,
              f"manuscript file has no {label} statement (Snapp double-anonymous instructions)")
    check("Author contributions" not in anon.split("## References")[0] and "Yusen" not in anon, "no author names or contributions in the anonymised manuscript")
    ors_cited = sorted({int(n) for n in re.findall(r"Online Resources? (\d+)", anon)} | {int(n) for n in re.findall(r"Online Resources \d+[–-](\d+)", anon)})
    listed = [int(n) for n in re.findall(r"^\*\*Online Resource (\d+)\*\*", anon, re.M)]
    check(bool(listed) and listed == list(range(1, len(listed) + 1)) and set(ors_cited) <= set(listed),
          "Online Resources: cited ones are listed with a caption, numbered 1..n", f"listed {listed}, cited {ors_cited}")
    main_part = anon.split("## Supplementary Information")[0]
    seen = []
    for m in re.finditer(r"Online Resources? (\d+)(?: and (\d+))?", main_part):
        for g in m.groups():
            if g and int(g) not in seen:
                seen.append(int(g))
    check(bool(listed) and seen == sorted(seen) and set(seen) == set(listed),
          "every Online Resource is cited in the text, first citations in numerical order", f"first citations: {seen}")
    sd_ors = {int(n) for n in re.findall(r"Online Resources? (\d+)", sd)} | {int(n) for n in re.findall(r"Online Resources \d+[–-](\d+)", sd)}
    check(bool(re.search(r"ESM_\d+\.(pdf|xlsx|zip|csv)", anon)) and all(re.search(rf"^\*\*Online Resource {i}\*\* \(ESM_{i}\.", anon, re.M) for i in listed),
          "each Online Resource is listed with its file name ESM_n.ext and a caption")
    check("⟦" not in anon, "no unresolved markers in the anonymised manuscript")
    left = re.findall(r"\[TO BE (?:SUPPLIED|CONFIRMED)[^\]]*\]", anon)
    check(not left, "no placeholders in the anonymised manuscript", "; ".join(left[:5]))
    todo = re.findall(r"\[TO BE (?:SUPPLIED|CONFIRMED)[^\]]*\]", page + letter)
    check(not todo, "title page and cover letter complete", f"{len(todo)} placeholders left", info=bool(todo))
    for label in ("Corresponding author", "Acknowledgements", "Author contributions", "Funding", "Competing interests", "Data availability"):
        check(f"**{label}**" in page, f"title page has: {label}")

    # ---- anonymity ------------------------------------------------------------------------------------------------
    for bad in REVEALING:
        check(bad not in anon, f"anonymised manuscript does not contain '{bad}'")

    # ---- Word files ------------------------------------------------------------------------------------------------------
    from docx import Document
    from docx.oxml.ns import qn
    for name in ("Manuscript_anonymised.docx", "Title_page.docx", "Cover_letter.docx"):
        p = SUB / name
        check(p.exists(), f"{name} exists")
        if not p.exists():
            continue
        z = zipfile.ZipFile(p)
        d = Document(str(p))
        sec = d.sections[0]
        check(abs(sec.page_width.mm - 210) < 1 and abs(sec.page_height.mm - 297) < 1, f"{name}: A4")
        footer = "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/footer"))
        doc_xml = z.read("word/document.xml").decode("utf-8")
        check(footer.count("PAGE") == 1 and "fldChar" not in doc_xml and "fldSimple" not in doc_xml,
              f"{name}: automatic page numbers, no other field codes")
        check("word/endnotes.xml" not in z.namelist(), f"{name}: no endnotes")
        check(len(re.findall(r"<w:(?:ins|del)\b", doc_xml)) == 0 and "<w:comment " not in z.read("word/comments.xml").decode("utf-8")
              if "word/comments.xml" in z.namelist() else len(re.findall(r"<w:(?:ins|del)\b", doc_xml)) == 0,
              f"{name}: no tracked changes or comments")
        normal = d.styles["Normal"]
        check("Times New Roman" in z.read("word/styles.xml").decode("utf-8"), f"{name}: Times New Roman")
        blob = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith((".xml", ".rels")))
        if name == "Manuscript_anonymised.docx":
            for bad in REVEALING:
                check(bad not in blob, f"{name}: no '{bad}' in any part of the file")
            check(d.core_properties.author == "" and d.core_properties.last_modified_by == "", f"{name}: author properties empty")
            check(len(d.tables) == len(caps), f"{name}: {len(d.tables)} tables as Word tables")
            check(all(t._tbl.tblPr.find(qn("w:tblBorders")) is not None for t in d.tables), f"{name}: three-line tables")
            check(all(all(r.bold for c in t.rows[0].cells for pp in c.paragraphs for r in pp.runs) for t in d.tables), f"{name}: bold header rows")
            check(len(d.inline_shapes) == 1 and abs(d.inline_shapes[0].width.mm - 119) < 1, f"{name}: one figure, 119 mm wide")
            descr = re.findall(r'<wp:docPr[^>]*descr="([^"]+)"', doc_xml)
            check(bool(descr), f"{name}: figure has alternative text")
            fn_xml = z.read("word/footnotes.xml").decode("utf-8") if "word/footnotes.xml" in z.namelist() else ""
            n_fn = len(re.findall(r"<w:footnote\b(?![^>]*w:type=)", fn_xml))
            check(True, f"{name}: footnotes", f"{n_fn} (the paper uses none; Springer asks for footnotes, never endnotes)", info=True)

    # ---- figure files -------------------------------------------------------------------------------------------------------
    from PIL import Image
    for name in ("Fig1.png", "Fig1.eps", "Fig1.tif"):
        check((SUB / name).exists(), f"{name} exists (Springer naming: Fig + number)")
    if (SUB / "Fig1.png").exists():
        im = Image.open(SUB / "Fig1.png")
        w, h = im.size
        check(w >= 1500 and h >= 1200, "Snapp: figure at least 1500 x 1200 pixels", f"{w} x {h}")
        check(im.info.get("dpi", (0, 0))[0] >= 599, "Fig1.png at 600 dpi", f"{im.info.get('dpi')}")
        check(abs(w / im.info["dpi"][0] * 25.4 - 119) < 1 and h / im.info["dpi"][1] * 25.4 <= 195, "figure 119 mm wide, at most 195 mm high",
              f"{w / im.info['dpi'][0] * 25.4:.0f} x {h / im.info['dpi'][1] * 25.4:.0f} mm")
        check(im.mode == "RGB", "RGB (8 bits per channel)")
        check(set(im.info) <= {"dpi"}, "Fig1.png: no metadata beyond the resolution (no author, no software, no path)", f"{sorted(im.info)}")
    if (SUB / "Fig1.tif").exists():
        t = Image.open(SUB / "Fig1.tif")
        check(t.info.get("dpi", (0, 0))[0] >= 599 and t.mode == "RGB", "Fig1.tif: 600 dpi RGB", f"{t.size}, {t.info.get('compression')}")
        check(not ({270, 271, 272, 305, 306, 315, 316, 33432} & set(t.tag_v2.keys())), "Fig1.tif: no description, camera, software, date, artist, host or copyright tags")
    if (SUB / "Fig1.eps").exists():
        eps = (SUB / "Fig1.eps").read_bytes()
        bb = re.search(rb"%%BoundingBox: 0 0 (\d+) (\d+)", eps)
        check(eps.startswith(b"%!PS-Adobe") and b"EPSF" in eps[:40], "Fig1.eps is an EPS file")
        check(bool(bb) and abs(int(bb.group(1)) / 72 * 25.4 - 119) < 1.5, "Fig1.eps 119 mm wide", f"{int(bb.group(1)) / 72 * 25.4:.0f} mm" if bb else "")
        check(b"/sfnts" in eps or b"BeginFont" in eps, "Fig1.eps: fonts embedded", info=False)
        head = eps[:2000].decode("latin-1")
        check(not re.search(r"%%(For|Author|Copyright):", head) and not any(b.encode() in eps for b in REVEALING), "Fig1.eps: no author, for or copyright line, no author's name")

    # ---- English as the main language ---------------------------------------------------------------------------------------------
    para_cjk = [p for p in re.split(r"\n\n+", anon.split("## References")[0]) if not p.startswith("|")
                and len(re.findall(r"[㐀-鿿]", p)) > len(re.findall(r"[A-Za-z]", p))]
    check(not para_cjk, "no paragraph that is mostly Chinese", f"{len(para_cjk)} found")
    allcjk = len(re.findall(r"[㐀-鿿\U00020000-\U0002fa1f]", anon.split("## References")[0]))
    allw = len(re.findall(r"[A-Za-z]+", anon.split("## References")[0]))
    check(True, "Chinese characters in the text", f"{allcjk} characters against {allw} Latin-script words; they are examples and data, with pinyin or English", info=True)

    # ---- supplement (Online Resources 1-7) ---------------------------------------------------------------------------------------
    supplement_checks(anon, rerun)

    width = max(len(n) for _, n, _ in results)
    for status, name, detail in results:
        print(f"{status:4s}  {name:{min(width, 78)}s}  {detail}")
    n_fail = sum(1 for s, _, _ in results if s == "FAIL")
    print(f"\n{sum(1 for s, _, _ in results if s == 'PASS')} passed, {n_fail} failed, {sum(1 for s, _, _ in results if s == 'INFO')} for information")
    return 1 if n_fail else 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    sys.exit(main(args[0] if args else "v10", rerun="--rerun-esm7" in sys.argv))
