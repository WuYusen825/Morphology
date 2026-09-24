# 亦声论文用的盲编码表：从两端都有 BS 构拟的条目中抽 40 个亦声字 + 60 个普通形声字，打乱顺序；
# 只给声符字和成员字的释义（去掉"从某某聲/亦聲"的结构分析），编码人看不到类别。
import json,random,re,csv,sys,os
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font,Alignment
from coding_yisheng import C
out=sys.argv[1]
R=json.load(open('yisheng_rows.json'))
ok=[r for r in R if r['morph_relation_auto']!='NA' and r['relation'] in ('亦聲','聲')]
rnd=random.Random(20260924)
yi=[r for r in ok if r['relation']=='亦聲']; pl=[r for r in ok if r['relation']=='聲']
samp=rnd.sample(yi,40)+rnd.sample(pl,60); rnd.shuffle(samp)
def defn(g):
    d=g.split('从')[0]; d=re.sub(r'凡.{1,3}之屬皆$','',d).strip()
    return d or g[:40]
wb=Workbook(); ws=wb.active; ws.title='编码'
ws.append(['item','声符字','声符字释义','成员字','成员字释义','语义关系(Y/E/N/X)','把握(1-3)','备注'])
for i,r in enumerate(samp,1): ws.append([i,r['phonetic'],defn(r['head_gloss']),r['char'],defn(r['shuowen_gloss']),'','',''])
n=len(samp)+1
dv=DataValidation(type='list',formula1='"Y,E,N,X"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'F2:F{n}')
dv2=DataValidation(type='list',formula1='"1,2,3"',allow_blank=True); ws.add_data_validation(dv2); dv2.add(f'G2:G{n}')
for col,w in zip('ABCDEFGH',[6,8,40,8,40,14,9,24]): ws.column_dimensions[col].width=w
for c in ws[1]: c.font=Font(bold=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
ws.freeze_panes='A2'
g=wb.create_sheet('说明',0)
for line in ['盲编码说明（亦声论文用；请勿查看 yisheng_dataset.csv 和 coding_yisheng.py）','',
 '任务：比较“成员字”的《说文》本义与“声符字”的《说文》本义，判断两者的语义关系。',
 'Y = 相同或近义（如 厓“山邊也” / 涯“水邊也”）',
 'E = 经一步引申或推理才能连上（如 亡“逃也” / 忘“不識也”）',
 'N = 看不出语义关系',
 'X = 成员字是专名（地名、水名、姓氏、动植物名等），无法判断',
 '把握：1 = 不确定，2 = 较确定，3 = 很确定',
 '','规则：只看本义，不看后起义和假借义。表中已去掉许慎“从某某声/亦声”的结构分析，请不要去查原书中这个字是不是亦声。',
 f'共 {len(samp)} 条，随机抽样，随机种子 20260924。']: g.append([line])
g.column_dimensions['A'].width=110
wb.save(os.path.join(out,'blind_coding_sheet.xlsx'))
with open(os.path.join(out,'yisheng_claude_codes.csv'),'w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(['item','series','char','relation','morph_relation_auto','code_core'])
    for i,r in enumerate(samp,1): w.writerow([i,r['phonetic'],r['char'],r['relation'],r['morph_relation_auto'],C[i]])
print(len(samp))
