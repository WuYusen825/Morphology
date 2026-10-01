"""Build ESM_2.xlsx, ESM_3.xlsx and ESM_6.xlsx from the CSV tables of the kit (internal; not shipped)."""
import sys, json, re
sys.path.insert(0, '/tmp/claude-0/-home-user-Morphology/7d9f6fd5-d830-551f-9a43-e3ebcc390a52/scratchpad/esm/work')
from esm_common import *
import pandas as pd
from openpyxl import load_workbook

dd = read_csv(KIT / 'data_dictionary.csv')
def wb_text(s):
    """Descriptions of the CSV dictionary refer to CSV files; in the workbooks they refer to the sheets."""
    s = s.replace('data/recension_collation.csv', 'the sheet "recension_collation"')
    return re.sub(r'\bpairs\.csv\b', 'the sheet "pairs" of Online Resource 2', s)


def dict_for(file, sheet, extra=None, rename=None):
    rows = []
    for _, r in dd[dd.file == file].iterrows():
        rows.append((sheet, (rename or {}).get(r.column, r.column), wb_text(r.description)))
    return rows + (extra or [])

CPDESC = 'Unicode code point(s) of the character, written U+XXXX (the Chinese character itself is in the column of the same name without "_codepoint")'
pairs = read_csv(KIT / 'data/pairs.csv')
pos = {s: i for i, s in enumerate(pairs.sw_id)}        # order of the dataset

built = {}   # sheet frames for verification: {file: {sheet: df}}


def save(wb, esm_no, frames):
    fn = CAPTIONS[esm_no][0]
    out = STAGE / fn
    wb.save(out)
    built[fn] = frames
    print(f'saved {out} ({out.stat().st_size} bytes): ' + ', '.join(f'{k} {v.shape}' for k, v in frames.items()))


# ================================================================================================ ESM_2
def build_esm2():
    wb = new_workbook('Online Resource 2: pair-level dataset')
    frames = {}
    frames['pairs'] = pairs
    coll = read_csv(KIT / 'data/recension_collation.csv')
    frames['recension_collation'] = coll

    # register of the 104 print-edition checks
    sp = read_csv(Path('/home/user/Morphology/youwen/manuscript/daxu_spotcheck_v6.csv'))
    sp = sp.rename(columns={'字头': 'headword', '扫描文件': 'scan_file', 'PDF页码_从1开始': 'pdf_page_1based', '核查记录状态': 'check_status', '扫描所见片段或备注': 'remark_zh', '电子文本说解': 'electronic_text_gloss'})
    GENERIC = '扫描页已目验；该条未另存文字摘录。'
    SPECIAL = {
        '1437': ('', 'The headword, the gloss 在口所以言也別味也 and the formula 干亦聲 are all seen on the page.'),
        '3365': ('艸也；楚謂之葍，秦謂之藑；从舛，舛亦聲', 'The entry continues on PDF page 89.'),
        '8024': ('傷擊也；从手毀，毀亦聲', "The character 毀 is shown in the print edition's own form."),
        '3261': ('', 'The electronic text reads 「日月合宿从辰」; the print edition clearly reads 「日月合宿爲辰」; the following 「从會从辰，辰亦聲」 agrees.'),
    }
    rows = []
    for _, r in sp.iterrows():
        t = r.remark_zh
        if r.sw_id in SPECIAL:
            exc, rem = SPECIAL[r.sw_id]
        elif t == GENERIC:
            exc, rem = '', 'Scan page inspected visually; no text excerpt of this entry was saved.'
        else:
            exc, rem = t, 'The excerpt agrees with the electronic text.'
        rows.append(dict(sw_id=r.sw_id, headword=r.headword, headword_codepoint=cp(r.headword), scan_file=r.scan_file, pdf_page_1based=r.pdf_page_1based,
                         check_status=r.check_status, scan_excerpt=exc, remark=rem, electronic_text_gloss=r.electronic_text_gloss))
    spot = pd.DataFrame(rows)
    assert len(spot) == 104 and spot.sw_id.is_unique
    assert set(spot.sw_id) <= set(pairs.sw_id)
    assert spot.check_status.value_counts().to_dict() == {'match': 73, 'visually_checked': 30, 'variant': 1}
    frames['daxu_spotcheck'] = spot

    dict_rows = (dict_for('data/pairs.csv', 'pairs') + dict_for('data/recension_collation.csv', 'recension_collation') + [
        ('daxu_spotcheck', 'sw_id', 'row number of the head entry (as in the sheet "pairs")'),
        ('daxu_spotcheck', 'headword', 'the labelled compound whose entry was checked'),
        ('daxu_spotcheck', 'headword_codepoint', CPDESC),
        ('daxu_spotcheck', 'scan_file', "file name of the scan (Waseda University Library, Chen Changzhi's 1873 print of the Shuōwén; ten PDF files numbered 0001-0010)"),
        ('daxu_spotcheck', 'pdf_page_1based', 'page of the PDF file, counted from 1 (not the leaf number of the print)'),
        ('daxu_spotcheck', 'check_status', 'match: the passage on the scan agrees with the electronic text (an excerpt was recorded); visually_checked: the page was inspected, no excerpt saved; variant: one wording differs (see remark)'),
        ('daxu_spotcheck', 'scan_excerpt', 'fragment of the print edition as recorded when the page was read (Chinese data); empty where none was saved'),
        ('daxu_spotcheck', 'remark', 'English remark on the check'),
        ('daxu_spotcheck', 'electronic_text_gloss', 'the Shuōwén gloss of the entry in the electronic text, including the structural analysis (Chinese data)')])
    dictdf = pd.DataFrame(dict_rows, columns=['sheet', 'column', 'description'])
    # every column of every data sheet must be described
    for sh in ['pairs', 'recension_collation', 'daxu_spotcheck']:
        assert set(dictdf[dictdf.sheet == sh].column) == set(frames[sh].columns), sh
    frames['data_dictionary'] = dictdf

    sheet_rows = [('pairs', *pairs.shape, 'One row per pair of a Shuōwén compound and its phonetic, as extracted from the open Dà Xú text (1,333 pairs: 223 labelled, 1,054 ordinary, 17 with an abbreviated phonetic, 39 extra rows found only in Duan Yucai or the Xiǎo Xú); the analysis set is the 212 labelled and 953 ordinary rows with in_comparison_frame = True. Middle and Old Chinese readings, sound-relation categories, gloss and recension flags, relation types of homophonous pairs. Tables 2, 4, 5, 6, 9, 10, 11.'),
                  ('recension_collation', *coll.shape, 'The Dà Xú–Xiǎo Xú collation of the 227 raw yìshēng headwords of the Dà Xú text (140 labelled in both recensions, 62 only in the Dà Xú, 10 undecidable among the 212 analysis entries). Table 10, Section 4.6.'),
                  ('daxu_spotcheck', *spot.shape, 'Register of the 104 analysis-set entries checked against the scan of a print edition (not a random sample; see Online Resource 1, Section 6). Section 3.1 and Section 6 of the paper.'),
                  ('data_dictionary', *dictdf.shape, 'Description of every column of the three data sheets.')]
    notes = COMMON_NOTES + [
        'group: labelled = the Dà Xú analysis names the phonetic with yìshēng 亦聲; ordinary = plain phonetic compound 聲; the 212 labelled and 953 ordinary rows with in_comparison_frame = True form the analysis set.',
        'Sound-relation categories: mc_relation (identical / tone_voicing_alt / other) for Middle Chinese; morph_relation_auto (I, R, O, O2, V, C) for Old Chinese (Baxter and Sagart 2014); definitions in the sheet "data_dictionary" and in Online Resource 1, Section 4.',
        'identical_relation_type (L, F, C, U, X) was assigned with the label visible and is descriptive only (Online Resource 1, Sections 4 and 5).',
        'The register of print-edition checks covers 104 of the 212 analysis entries, chosen as the checking proceeded; it must not be used to estimate a transcription error rate for the others (Online Resource 1, Section 6).']
    add_about(wb, 2, [], sheet_rows, notes)
    for sh in ['pairs', 'recension_collation', 'daxu_spotcheck']:
        add_table_sheet(wb, sh, frames[sh])
    add_table_sheet(wb, 'data_dictionary', dictdf, text_cols=('sheet', 'column', 'description'))
    save(wb, 2, frames)


# ================================================================================================ ESM_3
def build_esm3():
    wb = new_workbook('Online Resource 3: semantic codings')
    frames = {}
    fs = read_csv(KIT / 'data/first_sample_codes.csv')
    frames['first_sample'] = fs

    codes = read_csv(KIT / 'output/ext_codes_long.csv')
    ans = read_csv(KIT / 'output/ext_items_with_answers.csv')
    key = read_csv(KIT / 'data/ext_items_key.csv')
    assert list(ans.sw_id) == list(key.sw_id)
    A = ans.set_index('sw_id')
    FS = fs[fs.sw_id != ''].set_index('sw_id')
    anchors = set(key[key.role == 'anchor'].sw_id)
    frame = pairs[pairs.in_comparison_frame == 'True']
    assert frame.sw_id.is_unique and len(frame) == 1165
    P = frame.set_index('sw_id')
    rows = []
    for _, r in codes.iterrows():
        sid = r.sw_id
        new = r.source == 'extended'
        role = 'new' if new else ('anchor' if sid in anchors else 'first_sample')
        has_mc = P.loc[sid, 'mc_relation'] not in ('NA', '')
        has_oc = P.loc[sid, 'morph_relation_auto'] not in ('NA', '')
        d = dict(sw_id=sid, phonetic=r.phonetic, phonetic_codepoint=cp(r.phonetic), char=r.char, char_codepoint=cp(r.char),
                 group='labelled' if r.label == '1' else 'ordinary', relation=r.relation, source=r.source, role=role,
                 stratum='labelled' if r.label == '1' else ('ordinary_mc_oc' if has_oc else 'ordinary_mc_only'), has_mc=str(has_mc), has_oc=str(has_oc))
        if role in ('new', 'anchor'):
            a = A.loc[sid]
            d.update(phonetic_gloss=a.phonetic_gloss, member_gloss=a.member_gloss, pass1_batch=a.p1_batch, pass1_item=a.p1_item, pass2_batch=a.p2_batch, pass2_item=a.p2_item)
        else:
            f = FS.loc[sid]
            d.update(phonetic_gloss=f.phonetic_gloss, member_gloss=f.member_gloss, pass1_batch='', pass1_item='', pass2_batch='', pass2_item='')
        if new:
            a = A.loc[sid]
            assert (a.p1_code, a.p2_code, a.p1_conf, a.p2_conf) == (r.code1, r.code2, r.conf1, r.conf2)
            d.update(code1=a.p1_code, conf1=a.p1_conf, note1=a.p1_note, code2=a.p2_code, conf2=a.p2_conf, note2=a.p2_note)
        else:
            f = FS.loc[sid]
            assert (f.blind_code, f.blind_confidence) == (r.code1, r.conf1) == (r.code2, r.conf2)
            d.update(code1=f.blind_code, conf1=f.blind_confidence, note1=f.blind_note, code2=f.blind_code, conf2=f.blind_confidence, note2=f.blind_note)
        if role == 'anchor':
            a = A.loc[sid]
            d.update(anchor_pass1_code=a.p1_code, anchor_pass1_conf=a.p1_conf, anchor_pass2_code=a.p2_code, anchor_pass2_conf=a.p2_conf)
        else:
            d.update(anchor_pass1_code='', anchor_pass1_conf='', anchor_pass2_code='', anchor_pass2_conf='')
        d.update(cat4=r.cat4, paronomastic_gloss=r.paronomastic_gloss, layer=r.layer, mc_relation=r.mc_relation, label=r.label, near=r.near, ident=r.ident,
                 in_tests=str(r.mc_relation != '' and r.code1 != 'X'))
        rows.append(d)
    enl = pd.DataFrame(rows)
    enl = enl.iloc[sorted(range(len(enl)), key=lambda i: pos[enl.sw_id.iloc[i]])].reset_index(drop=True)
    assert len(enl) == 767 and enl.sw_id.is_unique
    assert enl.role.value_counts().to_dict() == {'new': 677, 'first_sample': 55, 'anchor': 35}
    assert (enl.in_tests == 'True').sum() == 634 and ((enl.in_tests == 'True') & (enl.group == 'labelled')).sum() == 172
    frames['enlarged_coding'] = enl

    ac = read_csv(KIT / 'data/author_check.csv')
    ck = read_csv(KIT / 'output/ext_check_key.csv').set_index('check_id')
    ac['sw_id'] = [ck.loc[c, 'sw_id'] for c in ac.check_id]
    ac['group'] = ['labelled' if ck.loc[c, 'relation'] == '亦聲' else 'ordinary' for c in ac.check_id]
    assert all(ck.loc[c, 'p1_code'] == p for c, p in zip(ac.check_id, ac.pass1_code))
    frames['author_check'] = ac

    sa = read_csv(KIT / 'output/ext_second_answer_items.csv')
    K = key.set_index('sw_id')
    sa2 = pd.DataFrame(dict(pass_no='1', batch=sa.batch, item=sa['item'], sw_id=sa.sw_id, phonetic=[K.loc[s, 'phonetic'] for s in sa.sw_id], char=[K.loc[s, 'char'] for s in sa.sw_id],
                            role=sa.role, first_code=sa['first'], first_conf=sa.conf1, second_code=sa['second'], second_conf=sa.conf2))
    assert len(sa2) == 203
    frames['second_answers'] = sa2

    dict_rows = (dict_for('data/first_sample_codes.csv', 'first_sample') + [
        ('enlarged_coding', 'sw_id', 'row number in the sheet "pairs" of Online Resource 2'),
        ('enlarged_coding', 'phonetic', 'phonetic of the pair'), ('enlarged_coding', 'phonetic_codepoint', CPDESC),
        ('enlarged_coding', 'char', 'compound of the pair'), ('enlarged_coding', 'char_codepoint', CPDESC),
        ('enlarged_coding', 'group', 'labelled (the analysis names the phonetic with yìshēng) / ordinary (plain phonetic compound); never shown to the coders'),
        ('enlarged_coding', 'relation', 'the same in the Shuōwén\'s own terms: 亦聲 labelled, 聲 ordinary'),
        ('enlarged_coding', 'source', 'extended: one of the 677 new items; original_100: one of the 90 first-sample items that lie in the sampling frame'),
        ('enlarged_coding', 'role', 'new: coded for the first time in passes 1 and 2; anchor: first-sample item added to a batch as a calibration check (5 per batch, 35 items in all, each coded once in each pass); first_sample: first-sample item in the frame that was not re-coded (55 items)'),
        ('enlarged_coding', 'stratum', 'sampling stratum: labelled / ordinary_mc_oc (ordinary, with Old Chinese forms at both ends) / ordinary_mc_only'),
        ('enlarged_coding', 'has_mc', 'True if both ends have Middle Chinese readings'), ('enlarged_coding', 'has_oc', 'True if both ends have Old Chinese reconstructions'),
        ('enlarged_coding', 'phonetic_gloss', 'Shuōwén gloss of the phonetic as shown to the coders (structural analysis removed); empty where the phonetic is not a head entry of the Shuōwén data (8 phonetics, 10 pairs), for which the coders were shown the word None'),
        ('enlarged_coding', 'member_gloss', 'Shuōwén gloss of the compound as shown to the coders (structural analysis removed)'),
        ('enlarged_coding', 'pass1_batch', 'batch (1-7) of the item in pass 1; empty for first_sample items'), ('enlarged_coding', 'pass1_item', 'item number within the batch in pass 1; empty for first_sample items'),
        ('enlarged_coding', 'pass2_batch', 'batch (1-7) of the item in pass 2; empty for first_sample items'), ('enlarged_coding', 'pass2_item', 'item number within the batch in pass 2; empty for first_sample items'),
        ('enlarged_coding', 'code1', 'code analysed: the pass-1 code for new items; for first-sample items the blind code of the first sample (Y same or near-synonymous meaning; E connected by one step of extension or inference; N no visible relation; X proper name)'),
        ('enlarged_coding', 'conf1', 'confidence of code1, 1 (low) to 3 (high)'), ('enlarged_coding', 'note1', "the coder's note for code1"),
        ('enlarged_coding', 'code2', 'second code: the pass-2 code for new items; for first-sample items the same blind code as code1'),
        ('enlarged_coding', 'conf2', 'confidence of code2; for first-sample items (anchors and first_sample) the same as conf1'), ('enlarged_coding', 'note2', "the coder's note for code2; for first-sample items the same as note1"),
        ('enlarged_coding', 'anchor_pass1_code', 'for the 35 anchors only: the code the pass-1 instance returned (compare with code1, the first-sample blind code)'),
        ('enlarged_coding', 'anchor_pass1_conf', 'confidence of that code'),
        ('enlarged_coding', 'anchor_pass2_code', 'for the 35 anchors only: the code the pass-2 instance returned'), ('enlarged_coding', 'anchor_pass2_conf', 'confidence of that code'),
        ('enlarged_coding', 'cat4', 'Middle Chinese category: identical / alt_departing (tone or voicing alternation involving the departing tone) / alt_other (other tone or voicing alternation) / other; empty without readings'),
        ('enlarged_coding', 'paronomastic_gloss', 'True if the definitional part of the gloss contains the phonetic character itself'),
        ('enlarged_coding', 'layer', 'recension layer of a labelled pair: both / daxu_only / unknown; empty for ordinary pairs'),
        ('enlarged_coding', 'mc_relation', 'identical / tone_voicing_alt / other; empty without readings'),
        ('enlarged_coding', 'label', '1 labelled, 0 ordinary'), ('enlarged_coding', 'near', '1 if cat4 is identical or alt_departing ("near"), else 0'), ('enlarged_coding', 'ident', '1 if cat4 is identical, else 0'),
        ('enlarged_coding', 'in_tests', 'True if the pair enters the tests: Middle Chinese readings at both ends and code1 not X (634 pairs: 172 labelled, 462 ordinary)')] +
        dict_for('data/author_check.csv', 'author_check', extra=[
            ('author_check', 'sw_id', 'row number in the sheet "pairs" of Online Resource 2 (added for linking; not shown to the authors)'),
            ('author_check', 'group', 'labelled / ordinary (added for linking; not shown to the authors)')]) + [
        ('second_answers', 'pass_no', 'blind pass in which the second answer was returned (always 1)'),
        ('second_answers', 'batch', 'batch number (3 or 6)'), ('second_answers', 'item', 'item number within the batch'), ('second_answers', 'sw_id', 'row number in the sheet "pairs" of Online Resource 2'),
        ('second_answers', 'phonetic', 'phonetic of the pair'), ('second_answers', 'char', 'compound of the pair'),
        ('second_answers', 'role', 'new / anchor'), ('second_answers', 'first_code', 'code of the first answer (the one that counted)'), ('second_answers', 'first_conf', 'confidence of the first answer'),
        ('second_answers', 'second_code', 'code of the second answer (not used in the main analysis)'), ('second_answers', 'second_conf', 'confidence of the second answer')])
    dictdf = pd.DataFrame(dict_rows, columns=['sheet', 'column', 'description'])
    for sh in ['first_sample', 'enlarged_coding', 'author_check', 'second_answers']:
        assert set(dictdf[dictdf.sheet == sh].column) == set(frames[sh].columns), (sh, set(dictdf[dictdf.sheet == sh].column) ^ set(frames[sh].columns))
    frames['data_dictionary'] = dictdf

    sheet_rows = [('first_sample', *fs.shape, 'First blind sample: 100 pairs (40 labelled, 60 ordinary) with the blind code, the non-blind code, confidence and notes. Section 3.3; Tables 3, 7, 11.'),
                  ('enlarged_coding', *enl.shape, 'Enlarged blind coding: the 767 coded pairs (677 new items coded in two blind passes, 90 first-sample items in the frame), with both passes\' codes, confidence and notes, the sampling stratum, the Middle Chinese category and the analysis flags. Section 3.4; Tables 2, 7, 8; Fig. 1.'),
                  ('author_check', *ac.shape, "The authors' consolidated check of 126 pairs (all 76 pairs on which the two passes disagree and 50 drawn at random from the 601 on which they agree). Section 4.3; Tables 7, 8."),
                  ('second_answers', *sa2.shape, 'The second answers returned by the Claude instances of pass 1, batches 3 and 6 (193 new items and 10 anchors), next to the first answers that were used. Table 8 note; Fig. 1.'),
                  ('data_dictionary', *dictdf.shape, 'Description of every column of the four data sheets.')]
    notes = COMMON_NOTES + [
        'Codes: Y same or near-synonymous meaning of the two characters; E connected only by one step of extension or inference; N no visible semantic relation; X the compound is a proper name. Y and E together count as "related". The coding instructions are in Online Resource 1, Section 4.',
        'The blind coders (Claude instances) never saw the group, the Shuōwén structural analysis or any file; the first sample was coded blind once (and once non-blind for comparison), the 677 new items twice in two independent passes. The authors\' check shows the consolidated judgment of the two authors, who checked separately (Online Resource 1, Section 5).',
        'The authors\' own working sheets for the first sample were not kept; the blind codes in the sheet "first_sample" are the codes of record (Online Resource 1, Section 5).']
    add_about(wb, 3, [], sheet_rows, notes)
    for sh in ['first_sample', 'enlarged_coding', 'author_check', 'second_answers']:
        add_table_sheet(wb, sh, frames[sh])
    add_table_sheet(wb, 'data_dictionary', dictdf, text_cols=('sheet', 'column', 'description'))
    save(wb, 3, frames)


# ================================================================================================ ESM_6
RESULTS = [  # (sheet, csv, content)
    ('summary_counts', 'yisheng_summary.csv', 'Counts and one-sided Fisher tests of the core comparisons of labelled and ordinary pairs (identity, alternations, departing-tone alternants, affix types, glosses, relation types, recension layers). Script core_models.py. Tables 4, 6, 9; Sections 4.1-4.2, 4.6.'),
    ('core_models', 'yisheng_models.csv', 'Mixed-effects (variational Bayes) and GEE estimates for H1a, H1b, H1c with Holm correction. Script core_models.py. Table 6; Sections 4.1-4.2.'),
    ('extra_checks', 'yisheng_models_extra_checks.csv', 'Further analyses of the paper: the four reading categories (Table 4), within-phonetic estimates, departing-tone steps (Table 6), recension layers (Table 10), combined checks (Table 11). Script paper_checks.py.'),
    ('wangyun_sensitivity', 'yisheng_models_wangyun_sensitivity.csv', "Sensitivity analyses without the nine labels rejected by Wang Yun. Script sensitivity_wang_yun.py. Table 11."),
    ('xiaoxu_sensitivity', 'yisheng_models_xiaoxu_sensitivity.csv', 'Sensitivity analyses by recension layer (labels shared by both recensions). Script sensitivity_xiao_xu.py. Tables 10 and 11.'),
    ('ext_coding_models', 'yisheng_models_ext_coding.csv', 'Enlarged blind coding: the primary test (H5), secondary tests S1-S9, exploratory tests X1-X4 and descriptive rows. Script ext_analysis.py. Tables 7 and 8; Fig. 1.'),
    ('ext_author_check_models', 'yisheng_models_ext_author_check.csv', "The same tests with the authors' consolidated judgments substituted for the 126 checked items. Script ext_author_check.py. Tables 7 and 8."),
    ('ext_second_answer_models', 'yisheng_models_ext_second_answer.csv', 'The same tests with the second answers of pass 1, batches 3 and 6 substituted (X5). Script ext_second_answer_sensitivity.py. Table 8; Fig. 1.'),
    ('bias_sensitivity_grid', 'yisheng_bias_sensitivity_grid.csv', 'Tipping-point grid: odds ratio corrected for assumed false-positive rates of the coding. Script ext_bias_sensitivity.py. Section 5.2.'),
    ('bias_reverse_scenario', 'yisheng_bias_reverse_scenario.csv', 'Reverse-direction bias scenario (labelled pairs coded unrelated merged into the related group). Script ext_bias_reverse_scenario.py. Not discussed in the text.'),
    ('table5_members', 'table5_departing_tone_members.csv', 'The 16 labelled pairs whose member is the departing-tone alternant of the phonetic. Script paper_counts.py. Table 5.'),
    ('fig1_values', 'fig1_h5_forest_values.csv', 'Values plotted in Fig. 1 (odds ratios and intervals of the test of H5 and its sensitivity analyses). Script paper_figure.py. Fig. 1.'),
]
REPORTS = [  # (sheet, txt, content)
    ('report_kappa', 'ext_kappa_output.txt', 'Agreement between the two blind passes on the 677 new items and with the first-sample codes on the 35 anchors (kappa). Script ext_analysis.py. Section 3.4.'),
    ('report_author_check', 'ext_author_check_output.txt', "Agreement between the authors' consolidated judgments and the two passes; sensitivity analysis with their codes. Script ext_author_check.py. Sections 3.4, 4.3."),
    ('report_first_sample', 'first_sample_checks_output.txt', 'First sample: agreement of the blind and non-blind codings (kappa 0.773), H2 on the first sample, identity read from fanqie strings. Script first_sample_checks.py. Sections 3.3, 4.3; Tables 7, 11.'),
    ('report_bias_sensitivity', 'ext_bias_sensitivity_output.txt', 'Bias sensitivity (tipping-point) analysis in words and numbers. Script ext_bias_sensitivity.py. Section 5.2.'),
    ('report_second_answer', 'ext_second_answer_sensitivity_output.txt', 'Agreement of first and second answers; tests with the second answers substituted. Script ext_second_answer_sensitivity.py. Table 8 note.'),
    ('report_reverse_scenario', 'ext_bias_reverse_scenario_output.txt', 'Reverse-direction bias scenario. Script ext_bias_reverse_scenario.py. Not discussed in the text.'),
    ('report_counts', 'paper_counts_output.txt', 'Counts quoted in the text and Tables 2, 3, 5 and 9, each checked against the data. Script paper_counts.py.'),
]


def build_esm6():
    wb = new_workbook('Online Resource 6: result tables')
    frames = {}
    dict_rows = []
    sheet_rows = []
    dfs = {}
    for sh, f, content in RESULTS:
        df = read_csv(KIT / 'output' / f)
        frames[sh] = df
        dfs[sh] = df
        dict_rows += dict_for('output/' + f, sh)
        sheet_rows.append((sh, *df.shape, content))
        assert set(c for s, c, _ in dict_rows if s == sh) == set(df.columns), sh
    reports = {}
    for sh, f, content in REPORTS:
        lines = (KIT / 'output' / f).read_text(encoding='utf-8').rstrip('\n').split('\n')
        reports[sh] = lines
        frames[sh] = pd.DataFrame({'line_no': [str(i) for i in range(1, len(lines) + 1)], 'text': lines})
        sheet_rows.append((sh, len(lines), 2, 'Text report, one line per row: ' + content))
        dict_rows += [(sh, 'line_no', 'line number'), (sh, 'text', 'the line of the report (English text; Chinese characters appear only as data)')]
    dictdf = pd.DataFrame(dict_rows, columns=['sheet', 'column', 'description'])
    frames['data_dictionary'] = dictdf
    sheet_rows.append(('data_dictionary', *dictdf.shape, 'Description of every column of every sheet.'))
    notes = COMMON_NOTES[:2] + [
        'Each result sheet is the file of the same content in Online Resource 7 (output/ and expected/), written by the script named in the sheet list. Numbers are shown with the number of digits at which the scripts print them.',
        'Estimates are odds ratios (OR) with 95% confidence intervals from logistic GEE clustered by phonetic (columns gee_*), Mantel-Haenszel within phonetic series (mh_*), and, in the core models, variational-Bayes mixed models (glmm_vb_*). Counts are written "hits" (outcome present) and "n" for the labelled group (group1, or labelled_) and the ordinary group (group0, or ordinary_). Definitions of every column are in the sheet "data_dictionary".',
        'Rows are in the order in which the scripts write them; the tier and test columns identify each analysis (for example P1 is the primary test of H5).',
        'The sheets report_* hold the text reports of the scripts (agreement statistics, bias analyses, counts) so that they can be read without running the code.']
    add_about(wb, 6, [], sheet_rows, notes)
    for sh, f, content in RESULTS:
        add_table_sheet(wb, sh, frames[sh])
    for sh, f, content in REPORTS:
        add_text_sheet(wb, sh, reports[sh])
    add_table_sheet(wb, 'data_dictionary', dictdf, text_cols=('sheet', 'column', 'description'))
    save(wb, 6, frames)


# ================================================================================================ verification
def verify(fn, frames):
    wb = load_workbook(STAGE / fn, data_only=False)
    p = wb.properties
    assert (p.creator or '') == '' and (p.lastModifiedBy or '') == '', (fn, p.creator, p.lastModifiedBy)
    assert wb.sheetnames[0] == 'About' and wb.sheetnames[1:] == list(frames), (wb.sheetnames, list(frames))
    bad = 0
    for sh, df in frames.items():
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        assert list(rows[0]) == list(df.columns), (fn, sh, rows[0][:5])
        assert len(rows) - 1 == len(df), (fn, sh, len(rows) - 1, len(df))
        for i, (row, exp) in enumerate(zip(rows[1:], df.itertuples(index=False))):
            for c, (v, e) in enumerate(zip(row, exp)):
                if v is None: ok = e == ''
                elif isinstance(v, bool): ok = False
                elif isinstance(v, (int, float)):
                    try: ok = abs(v - float(e)) <= 1e-12 * max(1, abs(v)) and (not NUM_INT.match(e) or str(v) == e)
                    except ValueError: ok = False
                else: ok = v == e
                if not ok:
                    bad += 1
                    if bad < 10: print('MISMATCH', fn, sh, i + 2, df.columns[c], repr(v), repr(e))
    print(f'verified {fn}: {len(frames)} sheets, {bad} mismatching cells')
    assert bad == 0


if __name__ == '__main__':
    build_esm2(); build_esm3(); build_esm6()
    for fn, fr in built.items():
        verify(fn, fr)
