"""Compare stage_v1 (delivered) with stage (rebuilt): what changed, cell by cell, file by file, line by line (internal)."""
import sys, zipfile, io, difflib, hashlib
from pathlib import Path
import pymupdf
from openpyxl import load_workbook
E = Path('/tmp/claude-0/-home-user-Morphology/7d9f6fd5-d830-551f-9a43-e3ebcc390a52/scratchpad/esm')
A, B = E / 'stage_v1', E / 'stage'
for f in sorted(p.name for p in A.iterdir()):
    a, b = (A / f).read_bytes(), (B / f).read_bytes()
    if a == b:
        print(f'{f}: IDENTICAL'); continue
    print(f'{f}: changed ({len(a)} -> {len(b)} bytes)')
    if f.endswith('.xlsx'):
        wa, wb_ = load_workbook(io.BytesIO(a)), load_workbook(io.BytesIO(b))
        assert wa.sheetnames == wb_.sheetnames, (wa.sheetnames, wb_.sheetnames)
        n = 0
        for sh in wa.sheetnames:
            ra, rb = list(wa[sh].iter_rows(values_only=True)), list(wb_[sh].iter_rows(values_only=True))
            if len(ra) != len(rb): print('   sheet', sh, 'rows', len(ra), len(rb))
            for i, (x, y) in enumerate(zip(ra, rb), 1):
                for j, (u, v) in enumerate(zip(x, y), 1):
                    if u != v:
                        n += 1
                        print(f'   {sh}!R{i}C{j}:\n      - {str(u)[:330]}\n      + {str(v)[:330]}')
        print('   changed cells:', n)
    elif f.endswith('.zip'):
        za, zb = zipfile.ZipFile(io.BytesIO(a)), zipfile.ZipFile(io.BytesIO(b))
        assert za.namelist() == zb.namelist(), set(za.namelist()) ^ set(zb.namelist())
        for nm in za.namelist():
            x, y = za.read(nm), zb.read(nm)
            if x != y:
                print('   changed file:', nm)
                if nm.endswith(('.txt', '.csv', '.py', '.md')):
                    xl, yl = x.decode('utf-8').split('\n'), y.decode('utf-8').split('\n')
                    for d in difflib.unified_diff(xl, yl, lineterm='', n=0):
                        if d.startswith(('---', '+++', '@@')): continue
                        print('     ', d[:360])
                else:
                    print('      (binary/xlsx)', len(x), len(y))
    elif f.endswith('.pdf'):
        da, db = pymupdf.open(stream=a, filetype='pdf'), pymupdf.open(stream=b, filetype='pdf')
        print('   pages', len(da), len(db))
        ta = [' '.join(p.get_text().split()) for p in da]
        tb = [' '.join(p.get_text().split()) for p in db]
        for i, (x, y) in enumerate(zip(ta, tb), 1):
            if x != y:
                sm = difflib.SequenceMatcher(None, x.split(' '), y.split(' '), autojunk=False)
                for tag, i1, i2, j1, j2 in sm.get_opcodes():
                    if tag != 'equal':
                        print(f'   p{i}: [{tag}] - {" ".join(x.split(" ")[i1:i2])[:300]}\n         + {" ".join(y.split(" ")[j1:j2])[:300]}')
