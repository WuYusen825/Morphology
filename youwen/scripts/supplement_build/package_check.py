"""Package-level checks of the seven Online Resources in stage/ (internal; not shipped).
Names, formats, file properties, anonymity, captions against the manuscript, encodings."""
import sys, io, re, zipfile, hashlib
sys.path.insert(0, '/tmp/claude-0/-home-user-Morphology/7d9f6fd5-d830-551f-9a43-e3ebcc390a52/scratchpad/esm/work')
from pathlib import Path
import pymupdf
from openpyxl import load_workbook
from esm_common import *

D = Path(sys.argv[1]) if len(sys.argv) > 1 else STAGE
fails = []; n = 0


def ok(c, msg):
    global n
    n += 1
    if not c:
        fails.append(msg); print('  FAIL:', msg)


EXPECT = [CAPTIONS[i][0] for i in range(1, 8)]
files = sorted(p.name for p in D.iterdir())
ok(files == EXPECT, f'file names: {files} != {EXPECT}')
for f in files:
    ok(bool(re.fullmatch(r'ESM_[1-7]\.(pdf|xlsx|zip)', f)), f'name {f}')
    ok((D / f).stat().st_size < 16 * 2**20, f'size {f}')

# patterns for personal data; assembled so that this file does not match itself
PAT = [r'\bQu\b', r'\bWu\b', 'Yu' + 'sen', 'Wei' + 'yi', 'wuyu' + 'sen', 'bf' + 'su', r'edu\.' + 'cn', r'2025\d{7}', 'Beijing Fo' + 'reign', 'Fo' + 'reign Studies', '吴雨' + '森',
       r'/home/', r'/tmp/', r'/mnt/', r'C:\\Users', 'github\\.com/Wu', 'claude\\.ai', 'session_0', r'[\w.+-]+@[\w-]+\.[\w.-]{2,}', 'preregist', '预注册', 'ORCID', r'orcid\.org',
       r'\bSnapp\b', 'SNAPP', 'Springer Nature']
RX = [re.compile(p, re.I if p not in (r'\bQu\b', r'\bWu\b') else 0) for p in PAT]


PATHPAT = {r'/home/', r'/tmp/', r'/mnt/', r'C:\\Users'}


def scan(name, text):
    hits = []
    for rx in RX:
        if name.endswith('self_check.py') and rx.pattern in PATHPAT: continue     # the self-check lists the path patterns it looks for
        m = rx.search(text)
        if m: hits.append((rx.pattern, text[max(0, m.start() - 30):m.end() + 30].replace('\n', ' ')))
    ok(not hits, f'personal-data/path pattern in {name}: {hits[:3]}')


def xlsx_checks(name, data):
    wb = load_workbook(io.BytesIO(data), read_only=False)
    p = wb.properties
    ok((p.creator or '') == '' and (p.lastModifiedBy or '') == '', f'{name}: creator/lastModifiedBy blank ({p.creator!r}, {p.lastModifiedBy!r})')
    ok(not (p.description or p.keywords or p.subject or p.category or p.identifier or p.language or p.revision), f'{name}: other properties blank')
    z = zipfile.ZipFile(io.BytesIO(data))
    for zi in z.namelist():
        ok('comments' not in zi and 'externalLink' not in zi and 'vba' not in zi.lower() and 'printerSettings' not in zi, f'{name}: no comments/links/macros ({zi})')
    for zi in ('docProps/core.xml', 'docProps/app.xml'):
        t = z.read(zi).decode('utf-8')
        scan(f'{name}:{zi}', t)
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    ok(not v.startswith('='), f'{name}: formula-like text in {ws.title}: {v[:30]}')
                    texts.append(v)
    scan(f'{name} cell text', '\n'.join(texts))
    return wb


def pdf_checks(name, path):
    d = pymupdf.open(path)
    m = d.metadata
    ok((m.get('author') or '') == '' and (m.get('creator') or '') == '' and (m.get('producer') or '') == '' and (m.get('subject') or '') == '' and (m.get('keywords') or '') == '',
       f'{name}: PDF metadata blank {m}')
    try:
        ok(not d.get_xml_metadata(), f'{name}: no XMP metadata')
    except Exception:
        pass
    text = '\n'.join(pg.get_text() for pg in d)
    scan(f'{name} text', text)
    ok('Withheld for double-anonymous review' in text, f'{name}: identification block with withheld authors')
    ok(TITLE.replace("'", "'") .split('?')[0] in text, f'{name}: article title present')
    ok(JOURNAL in text, f'{name}: journal name present')
    ok(CAPTIONS[int(name[4])][1][:60] in text.replace('\n', ' '), f'{name}: caption present')
    links = sum(len(pg.get_links()) for pg in d)
    ok(links == 0, f'{name}: no hyperlinks in the PDF ({links})')
    ok(not d.is_encrypted, f'{name}: not encrypted')
    ok(len(d.embfile_names()) == 0, f'{name}: no embedded files')
    return d


def zip_checks(name, path):
    z = zipfile.ZipFile(path)
    ok(z.comment == b'', f'{name}: zip comment empty')
    names = z.namelist()
    tops = {n_.split('/')[0] for n_ in names}
    ok(len(tops) == 1, f'{name}: single top folder {tops}')
    ok(all(not n_.startswith(('__MACOSX', '.')) and '/.' not in n_ and '__pycache__' not in n_ and not n_.endswith('.pyc') for n_ in names), f'{name}: no hidden files')
    ok(all(re.fullmatch(r'[A-Za-z0-9_./-]+', n_) for n_ in names), f'{name}: ASCII file names')
    for zi in z.infolist():
        ok(zi.extra == b'' and zi.comment == b'', f'{name}: no extra field / comment in {zi.filename}')
        ok(zi.date_time == (2026, 10, 1, 0, 0, 0), f'{name}: neutral timestamp {zi.filename} {zi.date_time}')
        data = z.read(zi.filename)
        ext = Path(zi.filename).suffix.lower()
        if ext in ('.csv', '.txt', '.py', '.md', '.tsv', '.sha256'):
            try:
                t = data.decode('utf-8')
            except UnicodeDecodeError:
                ok(False, f'{name}: not UTF-8 {zi.filename}'); continue
            ok('\r' not in t, f'{name}: LF line endings {zi.filename}')
            scan(f'{name}:{zi.filename}', t)
        elif ext == '.xlsx':
            xlsx_checks(f'{name}:{zi.filename}', data)
        elif ext == '.png':
            ok(b'tEXt' not in data[:2000] and b'iTXt' not in data[:2000], f'{name}: PNG without text chunks {zi.filename}')
    return z


for f in files:
    p = D / f
    if f.endswith('.xlsx'):
        wb = xlsx_checks(f, p.read_bytes())
        ok(wb.sheetnames[0] == 'About' and wb.sheetnames[-1] == 'data_dictionary', f'{f}: first sheet About, last data_dictionary ({wb.sheetnames[0]}, {wb.sheetnames[-1]})')
        ws = wb['About']
        about = {r[0].value: r[1].value for r in ws.iter_rows(min_row=1, max_row=8) if r[0].value}
        ok(about.get('Article title') == TITLE and about.get('Journal') == JOURNAL and about.get('Authors') == WITHHELD, f'{f}: About block identification')
        ok(about.get('Caption') == CAPTIONS[int(f[4])][1], f'{f}: About caption equals the manuscript caption')
    elif f.endswith('.pdf'):
        pdf_checks(f, p)
    else:
        zip_checks(f, p)

# the captions against the manuscript's Supplementary Information section
ms = Path('/mnt/project-files/youwen/v7_work/submission/source/yisheng_paper_v10_anonymised.md')
if ms.exists():
    t = ms.read_text(encoding='utf-8')
    sec = t.split('## Supplementary Information')[1].split('## ')[0]
    for i in range(1, 8):
        m = re.search(rf'\*\*Online Resource {i}\*\* \((ESM_{i}\.\w+)\) (.*)', sec)
        ok(m is not None, f'manuscript lists Online Resource {i}')
        if m:
            cap = re.sub(r'\*', '', m.group(2)).strip()
            ok(m.group(1) == CAPTIONS[i][0], f'manuscript file name for Online Resource {i}')
            ok(cap == CAPTIONS[i][1], f'caption {i} equals the manuscript text:\n   ms : {cap}\n   esm: {CAPTIONS[i][1]}')
    da = re.search(r'\*\*Data availability\*\* (.*)', t).group(1)
    ok('Online Resources 4 and 5' in da and 'Online Resource 1 is a guide to all files' in da, 'data availability statement refers to the Online Resources')
else:
    print('manuscript not found; captions not compared')

print(f'{n} checks, {len(fails)} failed')
sys.exit(1 if fails else 0)
