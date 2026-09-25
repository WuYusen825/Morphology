# H3 完整分层与"两本共有亦声"敏感性分析（2026-09-25），依据 yisheng_xiaoxu_collation_final.csv（scripts/xiaoxu_merge.py）
# 变体：original；只保留两本共有的亦声（去掉仅大徐与无从判断）；去掉仅大徐（保留无从判断）；仅大徐改记为"聲"
# 另报 H3：在亦声内部比较 both 与 daxu_only 两层（双侧，探索性）
# 用法：python3 xiaoxu_sensitivity.py <outdir> <blind_coding_sheet_llm_coded.xlsx> <yisheng_claude_codes.csv> <yisheng_xiaoxu_collation_final.csv>
import json,math,csv,sys,os,warnings
import pandas as pd, statsmodels.api as sm, statsmodels.formula.api as smf
from statsmodels.genmod.bayes_mixed_glm import BinomialBayesMixedGLM
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
out,sheet,ccodes,coll=sys.argv[1:5]
LAY={r['sw_id']:r['layer'] for r in csv.DictReader(open(coll,encoding='utf-8-sig'))}
R=json.load(open('yisheng_rows.json'))
lay=lambda r:LAY.get(str(r['sw_id']),'') if r['relation']=='亦聲' else ''
ws=load_workbook(sheet)['编码']
code={str(r[0]):r[5] for r in ws.iter_rows(min_row=2,values_only=True)}
S=[dict(r,code=code[r['item']]) for r in csv.DictReader(open(ccodes,encoding='utf-8-sig'))]
DS={(r['phonetic'],r['char'],r['relation']):r for r in R}
def keep(r,drop,recode):
    l=lay(r)
    if l in drop: return None
    if l in recode: return dict(r,relation='聲')
    return r
def variant(drop,recode):
    rows=[x for x in (keep(r,drop,recode) for r in R if r['in_analysis_set']=='True') if x]
    yph={r['phonetic'] for r in rows if r['relation']=='亦聲'}
    return [r for r in rows if r['phonetic'] in yph]
def h2rows(drop,recode):
    o=[]
    for s in S:
        d=DS.get((s['series'],s['char'],s['relation']))
        if not d or d['in_analysis_set']!='True' or s['code']=='X': continue
        d=keep(d,drop,recode)
        if not d: continue
        o.append(dict(y=int(s['code'] in 'YE'),label=int(d['relation']=='亦聲'),phon=d['phonetic'],par=d['paronomastic_gloss']))
    return pd.DataFrame(o)
def fit(df,vb=True,two=False):
    g=smf.gee('y ~ label',groups='phon',data=df,family=sm.families.Binomial(),cov_struct=sm.cov_struct.Exchangeable()).fit()
    b,se=g.params['label'],g.bse['label']; p=g.pvalues['label'] if two else (g.pvalues['label']/2 if b>0 else 1-g.pvalues['label']/2)
    res=dict(n=len(df),n_phonetics=df.phon.nunique(),labelled_hits=int(df[df.label==1].y.sum()),labelled_n=int((df.label==1).sum()),
             ordinary_hits=int(df[df.label==0].y.sum()),ordinary_n=int((df.label==0).sum()),
             gee_OR=round(math.exp(b),2),gee_CI95=f'{math.exp(b-1.96*se):.2f}-{math.exp(b+1.96*se):.2f}',**{('gee_p_two_sided' if two else 'gee_p_one_sided'):'%.2g'%p})
    if vb:
        v=BinomialBayesMixedGLM.from_formula('y ~ label',{'phon':'0 + C(phon)'},df).fit_vb()
        i=list(v.model.exog_names).index('label'); bm,sd=v.fe_mean[i],v.fe_sd[i]
        res.update(glmm_vb_OR=round(math.exp(bm),2),glmm_vb_CI95=f'{math.exp(bm-1.96*sd):.2f}-{math.exp(bm+1.96*sd):.2f}')
    return res
res=[]
for vname,drop,recode in [('original',(),()),('both recensions only (drop Da Xu-only and unknown)',('daxu_only','unknown'),()),
                          ('drop Da Xu-only (keep unknown)',('daxu_only',),()),('Da Xu-only recoded as 聲',(),('daxu_only',))]:
    A=variant(drop,recode); mc=[r for r in A if r['mc_relation']!='NA']
    df=pd.DataFrame(dict(y=[int(r['mc_relation']=='identical') for r in mc],label=[int(r['relation']=='亦聲') for r in mc],phon=[r['phonetic'] for r in mc]))
    res.append(dict(variant=vname,hypothesis='H1a',outcome='MC identical syllable',**fit(df)))
    alt=[r for r in mc if r['mc_relation']=='tone_voicing_alt']
    df=pd.DataFrame(dict(y=[int(r['mc_qusheng_direction']=='member') for r in alt],label=[int(r['relation']=='亦聲') for r in alt],phon=[r['phonetic'] for r in alt]))
    res.append(dict(variant=vname,hypothesis='H1c',outcome='member departing | tone/voicing alternation',**fit(df)))
    h=h2rows(drop,recode)
    res.append(dict(variant=vname,hypothesis='H2',outcome='semantic Y/E (blind LLM coding, X excluded)',**fit(h,vb=False)))
    res.append(dict(variant=vname,hypothesis='H2',outcome='same, paronomastic glosses removed',**fit(h[h.par=='False'],vb=False)))
# H3：亦声内部，both（label=1）对 daxu_only（label=0），双侧
Y=[r for r in R if r['in_analysis_set']=='True' and r['relation']=='亦聲' and lay(r) in('both','daxu_only')]
m=[r for r in Y if r['mc_relation']!='NA']
for name,rows,pred in [('MC identical syllable',m,lambda r:r['mc_relation']=='identical'),
                       ('MC tone/voicing alternation',m,lambda r:r['mc_relation']=='tone_voicing_alt'),
                       ('member departing | alternation',[r for r in m if r['mc_relation']=='tone_voicing_alt'],lambda r:r['mc_qusheng_direction']=='member')]:
    df=pd.DataFrame(dict(y=[int(pred(r)) for r in rows],label=[int(lay(r)=='both') for r in rows],phon=[r['phonetic'] for r in rows]))
    try: f=fit(df,vb=False,two=True)
    except Exception as e: f=dict(n=len(df),labelled_hits=int(df[df.label==1].y.sum()),labelled_n=int((df.label==1).sum()),ordinary_hits=int(df[df.label==0].y.sum()),ordinary_n=int((df.label==0).sum()),gee_OR='fit failed')
    res.append(dict(variant='H3 within 亦聲: both (labelled) vs Da Xu-only (ordinary)',hypothesis='H3',outcome=name,**f))
h=[]
for s in S:
    d=DS.get((s['series'],s['char'],s['relation']))
    if d and d['in_analysis_set']=='True' and s['code']!='X' and lay(d) in('both','daxu_only'): h.append(dict(y=int(s['code'] in 'YE'),label=int(lay(d)=='both'),phon=d['phonetic']))
h=pd.DataFrame(h); res.append(dict(variant='H3 within 亦聲: both (labelled) vs Da Xu-only (ordinary)',hypothesis='H3',outcome='semantic Y/E (blind LLM)',n=len(h),labelled_hits=int(h[h.label==1].y.sum()),labelled_n=int((h.label==1).sum()),ordinary_hits=int(h[h.label==0].y.sum()),ordinary_n=int((h.label==0).sum())))
with open(os.path.join(out,'yisheng_models_xiaoxu_sensitivity.csv'),'w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for r in res for k in r)));w.writeheader();w.writerows(res)
for r in res: print(r['variant'][:40],'|',r['hypothesis'],r['outcome'][:30],f"{r['labelled_hits']}/{r['labelled_n']} vs {r['ordinary_hits']}/{r['ordinary_n']}",r.get('gee_OR',''),r.get('gee_CI95',''),r.get('gee_p_one_sided',r.get('gee_p_two_sided','')),r.get('glmm_vb_OR',''))
