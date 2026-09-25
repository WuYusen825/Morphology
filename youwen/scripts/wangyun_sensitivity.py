# 敏感性分析：剔除王筠《說文釋例》认为是大徐误增的亦声标注后，重跑论文报告的 H1a、H1c、H2 模型（2026-09-25）
# 卷三"凡此類皆大徐誤增"9 字；卷八另驳 媄。出处见 /mnt/project-files/youwen/shili_pages/README.md E、F 节
# 两种处理：drop = 整行剔除；as_sheng = 把这些字改记为普通"聲"（王筠认为标注有误，字仍是同声符的形声字）
# 用法：python3 wangyun_sensitivity.py <outdir> <blind_coding_sheet_llm_coded.xlsx> <yisheng_claude_codes.csv>
import json,math,csv,sys,os,warnings
import pandas as pd, statsmodels.api as sm, statsmodels.formula.api as smf
from statsmodels.genmod.bayes_mixed_glm import BinomialBayesMixedGLM
from openpyxl import load_workbook
warnings.filterwarnings('ignore')
out,sheet,ccodes=sys.argv[1:4]
W9=set('貧愾恇娶婚姻婢緉坪'); W10=W9|{'媄'}
R=json.load(open('yisheng_rows.json'))
# H2 用的盲编码（独立 LLM 子代理），按条目号与 Claude 编码表对上声符和成员字
ws=load_workbook(sheet)['编码']
code={str(r[0]):r[5] for r in ws.iter_rows(min_row=2,values_only=True)}
S=[dict(r,code=code[r['item']]) for r in csv.DictReader(open(ccodes,encoding='utf-8-sig'))]
DS={(r['phonetic'],r['char'],r['relation']):r for r in R}   # 愷 在豈部（亦聲）和心部（聲）各有一条，须连关系一起匹配

def variant(excl,mode):
    rows=[]
    for r in R:
        if r['in_analysis_set']!='True': continue
        if r['relation']=='亦聲' and r['char'] in excl:
            if mode=='drop': continue
            r=dict(r,relation='聲')
        rows.append(r)
    yph={r['phonetic'] for r in rows if r['relation']=='亦聲'}
    return [r for r in rows if r['phonetic'] in yph]
def h2rows(excl,mode):
    out=[]
    for s in S:
        d=DS.get((s['series'],s['char'],s['relation']))
        if not d or d['in_analysis_set']!='True' or s['code']=='X': continue
        rel=d['relation']
        if rel=='亦聲' and s['char'] in excl:
            if mode=='drop': continue
            rel='聲'
        out.append(dict(y=int(s['code'] in 'YE'),label=int(rel=='亦聲'),phon=d['phonetic'],par=d['paronomastic_gloss']))
    return pd.DataFrame(out)
def fit(df,vb=True):
    g=smf.gee('y ~ label',groups='phon',data=df,family=sm.families.Binomial(),cov_struct=sm.cov_struct.Exchangeable()).fit()
    b,se=g.params['label'],g.bse['label']; p=g.pvalues['label']/2 if b>0 else 1-g.pvalues['label']/2
    res=dict(n=len(df),n_phonetics=df.phon.nunique(),labelled_hits=int(df[df.label==1].y.sum()),labelled_n=int((df.label==1).sum()),
             ordinary_hits=int(df[df.label==0].y.sum()),ordinary_n=int((df.label==0).sum()),
             gee_OR=round(math.exp(b),2),gee_CI95=f'{math.exp(b-1.96*se):.2f}-{math.exp(b+1.96*se):.2f}',gee_p_one_sided='%.2g'%p)
    if vb:
        v=BinomialBayesMixedGLM.from_formula('y ~ label',{'phon':'0 + C(phon)'},df).fit_vb()
        i=list(v.model.exog_names).index('label'); bm,sd=v.fe_mean[i],v.fe_sd[i]
        res.update(glmm_vb_OR=round(math.exp(bm),2),glmm_vb_CI95=f'{math.exp(bm-1.96*sd):.2f}-{math.exp(bm+1.96*sd):.2f}')
    return res
res=[]
for vname,excl,mode in [('original',set(),'drop'),('drop Wang 9',W9,'drop'),('drop Wang 10 (+媄)',W10,'drop'),
                        ('Wang 9 recoded as 聲',W9,'as_sheng'),('Wang 10 recoded as 聲',W10,'as_sheng')]:
    A=variant(excl,mode)
    mc=[r for r in A if r['mc_relation']!='NA']
    df=pd.DataFrame(dict(y=[int(r['mc_relation']=='identical') for r in mc],label=[int(r['relation']=='亦聲') for r in mc],phon=[r['phonetic'] for r in mc]))
    res.append(dict(variant=vname,hypothesis='H1a',outcome='MC identical syllable',**fit(df)))
    alt=[r for r in mc if r['mc_relation']=='tone_voicing_alt']
    df=pd.DataFrame(dict(y=[int(r['mc_qusheng_direction']=='member') for r in alt],label=[int(r['relation']=='亦聲') for r in alt],phon=[r['phonetic'] for r in alt]))
    res.append(dict(variant=vname,hypothesis='H1c',outcome='member departing | tone/voicing alternation',**fit(df)))
    h=h2rows(excl,mode)
    res.append(dict(variant=vname,hypothesis='H2',outcome='semantic Y/E (blind LLM coding, X excluded)',**fit(h,vb=False)))
    res.append(dict(variant=vname,hypothesis='H2',outcome='same, paronomastic glosses removed',**fit(h[h.par=='False'],vb=False)))
with open(os.path.join(out,'yisheng_models_wangyun_sensitivity.csv'),'w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for r in res for k in r)));w.writeheader();w.writerows(res)
for r in res: print(r['variant'],'|',r['hypothesis'],r['outcome'][:30],f"{r['labelled_hits']}/{r['labelled_n']} vs {r['ordinary_hits']}/{r['ordinary_n']}",r['gee_OR'],r['gee_CI95'],r['gee_p_one_sided'],r.get('glmm_vb_OR',''))
